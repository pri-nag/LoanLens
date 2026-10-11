"""Export the CSV data to one Excel workbook for reviewers:  data/lender_rules.xlsx

The CSVs stay the source of truth (the engine reads them); this workbook is regenerated from them.
Sheets:
  About             what each sheet is, colour legend
  Lender rules      one row per lender, one column per rule; cell colour = evidence tier,
                    click a cell to open its source, hover for date and estimation method
  Rules (long)      one row per lender x rule with source URL, date and estimation method
  Datapoints        input variables of the model
  Rule catalog      definition of every rule
  Sources           source register (URL, publisher, date read)
  Lenders           the 20 lenders and why they were chosen
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from _path import ROOT
from loan_engine.rules import SEGMENTS, _read, load_catalog, load_lenders, load_sources

OUT = ROOT / "data/lender_rules.xlsx"
TIER_FILL = {
    "lender_official": PatternFill("solid", fgColor="C6EFCE"),  # green
    "aggregator": PatternFill("solid", fgColor="FFEB9C"),       # yellow
    "estimate": PatternFill("solid", fgColor="F8CBAD"),         # orange
}
TIER_LABEL = {"lender_official": "Lender's own website", "aggregator": "Aggregator site (Paisabazaar, "
              "BankBazaar, Groww, Wishfin, GoCredit)", "estimate": "Estimate (method given on the row)"}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
NUMERIC = {"int": int, "float": float}


def typed(value: str, rule_type: str):
    """Write numbers as numbers so Excel can sort and filter them; everything else as text."""
    try:
        return NUMERIC[rule_type](value) if rule_type in NUMERIC and value != "" else value
    except ValueError:
        return value


def is_url(s: str) -> bool:
    return s.startswith(("http://", "https://"))


def style_sheet(ws, widths: dict[int, int] | None = None, freeze: str = "A2"):
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        if widths and col in widths:
            ws.column_dimensions[letter].width = widths[col]
        else:
            longest = max((len(str(c.value)) for c in ws[letter][:200] if c.value is not None), default=8)
            ws.column_dimensions[letter].width = min(max(10, longest + 2), 60)


def table_sheet(wb, title: str, rows: list[dict], link_cols=("url",)):
    ws = wb.create_sheet(title)
    header = list(rows[0])
    ws.append(header)
    for r in rows:
        ws.append([r[h] for h in header])
    for col, name in enumerate(header, 1):
        if name in link_cols:
            for cell in ws[get_column_letter(col)][1:]:
                if isinstance(cell.value, str) and is_url(cell.value):
                    cell.hyperlink, cell.font = cell.value, LINK_FONT
    style_sheet(ws)
    return ws


def main():
    load_lenders()  # raises if the long table fails validation
    catalog, sources = load_catalog(), load_sources()
    lenders = _read("lenders.csv")
    rows = _read("rules_long.csv")
    by_lender: dict[str, dict] = {}
    for r in rows:
        by_lender.setdefault(r["lender_id"], {})[r["rule"]] = r
    present = {r["rule"] for r in rows}
    rule_order = [k for base in catalog for k in [base] + [f"{base}@{s}" for s in SEGMENTS] if k == base or k in present]
    read_dates = sorted({s["accessed_on"] for s in sources.values()})

    wb = Workbook()

    # ---- About ----
    ws = wb.active
    ws.title = "About"
    n_est = sum(r["is_estimate"] == "true" for r in rows)
    about = [
        ["Personal loan eligibility rules: 20 Indian lenders"],
        [f"Generated from the CSV files in data/ by scripts/export_excel.py. Do not edit this workbook; "
         f"edit the CSVs and re-run the script."],
        [f"{len(lenders)} lenders, {len(rows)} rules ({n_est} estimated). Sources read on: {', '.join(read_dates)}."],
        [],
        ["Sheet", "What it shows"],
        ["Lender rules", "One row per lender, one column per rule. Cell colour = evidence tier. Click a value to "
                         "open its source; hover for the date read and the estimation method."],
        ["Rules (long)", "Same rules, one row per lender x rule, with source URL, date and estimation method. "
                         "Filter on is_estimate or evidence_tier."],
        ["Datapoints", "The model's input variables: unit, type, allowed values, where each comes from "
                       "(bureau / applicant / document)."],
        ["Rule catalog", "Definition of every rule: unit, type, which applicants, which model stage."],
        ["Sources", "Every source: URL, publisher, type, date read."],
        ["Lenders", "The 20 lenders, their segment, and why each was chosen."],
        [],
        ["Colour", "Evidence tier"],
    ]
    for line in about:
        ws.append(line)
    for tier, fill in TIER_FILL.items():
        ws.append([tier, TIER_LABEL[tier]])
        ws.cell(ws.max_row, 1).fill = fill
    ws["A1"].font = Font(bold=True, size=14)
    for row in (5, 13):
        for cell in ws[row]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=6):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # ---- Lender rules (one row per lender) ----
    ws = wb.create_sheet("Lender rules")
    ws.append(["lender_id", "lender_name", "lender_type", "rules", "estimated"] + rule_order)
    for col, rule in enumerate(rule_order, 6):
        base, _, seg = rule.partition("@")
        c = catalog[base]
        who = {"govt": "Govt/PSU employees only", "rel": "Applicants with a salary account at this lender only",
               "govt_rel": "Govt/PSU employees with a salary account at this lender only"}.get(seg, c["applies_to"])
        ws.cell(1, col).comment = Comment(f"{c['description']}\nUnit: {c['unit']}\nApplies to: {who}",
                                          "catalog", width=300, height=120)
    for l in lenders:
        rules = by_lender.get(l["lender_id"], {})
        ws.append([l["lender_id"], l["lender_name"], l["lender_type"], len(rules),
                   sum(r["is_estimate"] == "true" for r in rules.values())])
        row = ws.max_row
        for col, rule in enumerate(rule_order, 6):
            r = rules.get(rule)
            if r is None:
                continue
            s = sources[r["source_id"]]
            cell = ws.cell(row, col, typed(r["value"], catalog[rule.partition("@")[0]]["type"]))
            cell.fill = TIER_FILL[r["evidence_tier"]]
            note = [f"Tier: {r['evidence_tier']}", f"Source: {s['source_id']} ({s['publisher']})",
                    f"Read on: {s['accessed_on']}"]
            if r["estimation_method"]:
                note.append(f"Estimated: {r['estimation_method']}")
            if r["notes"]:
                note.append(f"Note: {r['notes']}")
            cell.comment = Comment("\n".join(note), "rules_long.csv", width=320, height=140)
            if is_url(s["url"]):
                cell.hyperlink = s["url"]
    style_sheet(ws, widths={1: 12, 2: 26, 3: 11, 4: 7, 5: 10}, freeze="C2")
    for col in range(6, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.row_dimensions[1].height = 60

    # ---- Rules (long), joined with the source register ----
    long_rows = []
    for r in rows:
        s = sources[r["source_id"]]
        long_rows.append({
            "lender_id": r["lender_id"], "rule": r["rule"],
            "value": typed(r["value"], catalog[r["rule"].partition("@")[0]]["type"]),
            "unit": catalog[r["rule"].partition("@")[0]]["unit"],
            "evidence_tier": r["evidence_tier"], "is_estimate": r["is_estimate"],
            "estimation_method": r["estimation_method"], "source_id": r["source_id"],
            "publisher": s["publisher"], "url": s["url"], "accessed_on": s["accessed_on"], "notes": r["notes"]})
    ws = table_sheet(wb, "Rules (long)", long_rows)
    tier_col = list(long_rows[0]).index("evidence_tier") + 1
    for (cell,) in ws.iter_rows(min_row=2, min_col=tier_col, max_col=tier_col):
        cell.fill = TIER_FILL[cell.value]

    table_sheet(wb, "Datapoints", _read("input_variables.csv"))
    table_sheet(wb, "Rule catalog", _read("rule_catalog.csv"))
    table_sheet(wb, "Sources", _read("sources.csv"))
    table_sheet(wb, "Lenders", _read("lenders.csv"))

    # fixed timestamp, so re-running without data changes gives the same document properties
    stamp = datetime.strptime(read_dates[-1], "%Y-%m-%d")
    wb.properties.created = wb.properties.modified = stamp
    wb.save(OUT)
    print(f"Wrote {OUT.relative_to(ROOT).as_posix()} ({len(lenders)} lenders, {len(rows)} rules, {len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
